from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "Guide_utilisation_QA.docx"

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
TEAL = RGBColor(0x0F, 0x6E, 0x56)
GRAY = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_run(run, *, size=11, bold=False, color=GRAY, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=26, bold=True, color=NAVY)


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=12, color=MUTED)


def add_h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, size=16 if level == 1 else 13, bold=True, color=NAVY if level == 1 else TEAL)


def add_p(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=11)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_li(doc, text, level=0):
    p = doc.add_paragraph(style="List Number" if level == 0 else "List Bullet")
    p.clear()
    r = p.add_run(text)
    set_run(r, size=11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    r = p.add_run(text)
    set_run(r, size=11)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run(r, size=10, name="Consolas", color=RGBColor(0x1A, 0x1A, 0x1A))
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("À retenir : ")
    set_run(r, size=11, bold=True, color=TEAL)
    r2 = p.add_run(text)
    set_run(r2, size=11)
    p.paragraph_format.space_after = Pt(10)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    add_title(doc, "Guide d’utilisation")
    add_subtitle(doc, "Framework QA — quality gate pour projets web")
    add_subtitle(doc, "Mode d’emploi pas à pas pour débutant")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Version 1.0  ·  Windows  ·  PowerShell")
    set_run(r, size=10, color=MUTED)

    add_h(doc, "1. C’est quoi ?", 1)
    add_p(
        doc,
        "Le framework QA est un outil en ligne de commande. Dans le dossier d’un projet "
        "(par exemple pilote-web), tu tapes qa init puis qa test.",
    )
    add_bullet(
        doc,
        "qa init prépare le projet : config, scripts npm, fichiers d’outils et specs "
        "d’accueil manquants. Il n’écrase jamais un fichier déjà présent.",
    )
    add_bullet(
        doc,
        "qa test relance la même préparation si besoin, exécute les contrôles "
        "(lint, tests, sécurité, accessibilité, performance…), puis ouvre un rapport HTML.",
    )
    add_p(
        doc,
        "Tu n’as pas besoin de relancer chaque outil à la main.",
    )

    add_h(doc, "2. Ce dont tu as besoin", 1)
    add_bullet(doc, "Windows, avec PowerShell (le terminal déjà utilisé).")
    add_bullet(doc, "Node.js et npm installés (dans le terminal : node -v puis npm -v).")
    add_bullet(doc, "Le dossier qa_framework (le code de l’outil).")
    add_bullet(doc, "Le dossier du projet à tester (ex. pilote-web), avec un package.json.")
    add_note(
        doc,
        "Sur cet ordinateur, npm peut échouer à cause d’un certificat. Avant les commandes npm, "
        "qa init ou qa test, exécute : $env:NODE_OPTIONS='--use-system-ca'",
    )

    add_h(doc, "3. Préparer la commande qa (une seule fois)", 1)
    add_p(doc, "Cette étape installe la commande qa sur ton PC. Tu ne la refais pas à chaque projet.")
    add_li(doc, "Ouvre PowerShell.")
    add_li(doc, "Va dans le dossier du framework :")
    add_code(doc, 'cd "C:\\Users\\GrâceAdjetey\\Documents\\Mes projets\\qa_framework"')
    add_li(doc, "Active le correctif certificat, puis installe et compile :")
    add_code(
        doc,
        "$env:NODE_OPTIONS='--use-system-ca'\nnpm install\nnpm run build\nnpm link",
    )
    add_li(doc, "Vérifie que la commande existe :")
    add_code(doc, "qa --help")
    add_p(
        doc,
        "Tu dois voir Usage: qa et les commandes init et test. Si qa n’est pas reconnu, "
        "ferme le terminal, rouvre-le, et réessaie. Après une modification du code du "
        "framework, refais seulement npm run build (pas besoin de npm link à chaque fois).",
    )

    add_h(doc, "4. Lancer les tests sur un projet existant", 1)
    add_p(
        doc,
        "Un projet existant (comme pilote-web) a déjà un package.json. S’il manque des configs "
        "ou des scripts, qa init les ajoute sans toucher à ce qui existe déjà.",
    )
    add_li(doc, "Va dans le dossier du projet, pas dans qa_framework :")
    add_code(doc, 'cd "C:\\Users\\GrâceAdjetey\\Documents\\Mes projets\\pilote-web"')
    add_li(doc, "Active le correctif certificat (si tu viens d’ouvrir le terminal) :")
    add_code(doc, "$env:NODE_OPTIONS='--use-system-ca'")
    add_li(doc, "(Optionnel) complète ce qui manque, puis lance la quality gate :")
    add_code(doc, "qa init\nqa test")
    add_p(doc, "Pendant l’attente, un spinner (⠋ ⠙ ⠹ …) indique le contrôle en cours.")
    add_p(doc, "À la fin tu vois :")
    add_bullet(doc, "une ligne par contrôle : ✔ réussi, ✖ échec bloquant, ⚠ avertissement")
    add_bullet(doc, "PASS (tout va bien) ou FAIL (au moins un contrôle obligatoire a échoué)")
    add_bullet(doc, "le rapport HTML qui s’ouvre dans le navigateur")
    add_bullet(doc, "un lien file://… dans le terminal (Ctrl + clic si besoin)")
    add_note(
        doc,
        "Le rapport est aussi enregistré ici : dossier-du-projet\\.qa\\report.html",
    )

    add_h(doc, "5. Préparer un nouveau projet", 1)
    add_p(
        doc,
        "Deux commandes suffisent après npm init -y : qa init puis qa test.",
    )
    add_p(
        doc,
        "qa init crée uniquement ce qui manque : qa.config.json, scripts npm, configs "
        "d’outils, specs d’accueil e2e / a11y / visual, lignes .gitignore. Il n’écrase "
        "jamais un fichier déjà présent.",
    )
    add_p(
        doc,
        "qa test refait la même préparation si besoin, puis exécute les contrôles.",
    )

    add_h(doc, "5.1. Créer le projet et l’initialiser", 2)
    add_code(
        doc,
        'cd "C:\\Users\\GrâceAdjetey\\Documents\\Mes projets"\n'
        "mkdir mon-projet\n"
        "cd mon-projet\n"
        "npm init -y\n"
        "$env:NODE_OPTIONS='--use-system-ca'\n"
        "qa init\n"
        "qa test",
    )
    add_p(
        doc,
        "Si c’est une application web, crée-la avec ton outil habituel (Vite, etc.), "
        "puis lance qa init à la racine. Les fichiers déjà là (tsconfig, eslint…) restent intacts.",
    )

    add_h(doc, "5.2. qa.config.json", 2)
    add_p(
        doc,
        "qa init crée ce fichier s’il n’existe pas. Tu peux ensuite l’éditer. "
        "Sans ce fichier, qa test s’arrête avec : qa.config.json introuvable.",
    )
    add_p(doc, "Signification des deux drapeaux :")
    add_bullet(doc, "enabled: true = le contrôle est lancé. false = il est ignoré.")
    add_bullet(
        doc,
        "required: true = un échec fait échouer toute la gate (FAIL). false = l’échec "
        "devient un avertissement (⚠), la gate peut quand même être PASS.",
    )
    add_note(
        doc,
        "Sur un projet tout neuf, tu peux désactiver e2e / visual / lighthouse dans "
        "qa.config.json tant que tu n’as pas d’app qui tourne sur le port 5173.",
    )

    add_h(doc, "5.3. Scripts npm", 2)
    add_p(
        doc,
        "qa init ajoute les scripts manquants dans package.json. Il ne remplace pas "
        "un script déjà défini.",
    )
    rows = [
        ("Contrôle", "Script npm attendu", "Qui le crée ?"),
        ("lint", "lint", "qa init (s’il manque)"),
        ("typecheck", "typecheck", "qa init (s’il manque)"),
        ("unit", "test:unit", "qa init (s’il manque)"),
        ("integration", "test:integration", "Toi (scénario métier)"),
        ("coverage", "test:coverage", "qa init (s’il manque)"),
        ("e2e", "test:e2e", "qa init (s’il manque)"),
        ("a11y", "test:a11y", "qa init (s’il manque)"),
        ("visual", "test:visual", "qa init (s’il manque)"),
        ("audit", "security:audit", "qa init (s’il manque)"),
        ("secrets", "security:secrets", "qa init (s’il manque)"),
        ("bundle", "perf:bundle", "qa init (s’il manque)"),
        ("lighthouse", "perf:lighthouse", "qa init (s’il manque)"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.rows[i].cells[j].text = cell
            for paragraph in table.rows[i].cells[j].paragraphs:
                for run in paragraph.runs:
                    set_run(run, size=10, bold=(i == 0), color=NAVY if i == 0 else GRAY)
    doc.add_paragraph()

    add_h(doc, "5.4. Ce que qa init / qa test créent tout seuls", 2)
    add_p(doc, "Uniquement si le contrôle est activé et que l’élément n’existe pas encore :")
    add_bullet(doc, "lint : eslint + eslint.config.js")
    add_bullet(doc, "typecheck : typescript + tsconfig.json")
    add_bullet(doc, "unit / coverage : vitest + vitest.config.ts")
    add_bullet(doc, "e2e : Playwright, playwright.config.ts, e2e/accueil.spec.ts, Chromium")
    add_bullet(doc, "a11y : @axe-core/playwright + e2e/a11y/accueil.a11y.spec.ts")
    add_bullet(doc, "visual : e2e/visual/accueil.visual.spec.ts")
    add_bullet(doc, "audit : script security:audit")
    add_bullet(doc, "secrets : secretlint + .secretlintrc.json + .secretlintignore")
    add_bullet(doc, "bundle : size-limit + .size-limit.json")
    add_bullet(doc, "lighthouse : paquets + perf-lighthouse.mjs")
    add_p(doc, "Ensuite, lance :")
    add_code(doc, "$env:NODE_OPTIONS='--use-system-ca'\nqa init\nqa test")

    add_h(doc, "6. Les fichiers de tests", 1)
    add_p(
        doc,
        "qa init pose des specs d’accueil génériques (la page se charge, a11y, snapshot). "
        "Les scénarios métier restent à toi.",
    )
    add_bullet(doc, "Unitaires : fichiers *.test.ts (souvent dans src/).")
    add_bullet(doc, "Intégration : fichiers dans tests/ — à créer toi-même.")
    add_bullet(doc, "E2E : fichiers e2e/*.spec.ts (Playwright). Spec d’accueil créée si absente.")
    add_bullet(doc, "Accessibilité : e2e/a11y/*.spec.ts — spec d’accueil créée si absente.")
    add_bullet(
        doc,
        "Visuel : e2e/visual/*.spec.ts. Spec d’accueil créée si absente. Au premier run local, "
        "Playwright crée l’image de référence (PNG). Il faut la commiter. Un vrai changement "
        "d’écran fera échouer le test jusqu’à npm run test:visual -- --update-snapshots "
        "si le changement est voulu.",
    )

    add_h(doc, "7. Lire le résultat", 1)
    add_p(doc, "Dans le terminal :")
    add_bullet(doc, "✔ = le contrôle a réussi.")
    add_bullet(doc, "✖ = échec d’un contrôle obligatoire → FAIL global.")
    add_bullet(doc, "⚠ = échec d’un contrôle non obligatoire → la gate peut rester PASS.")
    add_p(doc, "Dans le navigateur (rapport HTML) :")
    add_bullet(doc, "le verdict PASS ou FAIL en haut")
    add_bullet(doc, "la liste des contrôles avec durée et résumé")
    add_bullet(doc, "le détail des erreurs pour les échecs")
    add_p(doc, "Fichiers produits :")
    add_code(doc, "projet\\.qa\\report.html\nprojet\\.qa\\report.json")

    add_h(doc, "8. Commandes utiles", 1)
    add_code(doc, "qa --help")
    add_p(doc, "Aide générale.")
    add_code(doc, "qa init")
    add_p(doc, "Prépare le projet (config, scripts, fichiers manquants) sans lancer les tests.")
    add_code(doc, 'qa init --cwd "C:\\chemin\\vers\\un-autre-projet"')
    add_p(doc, "Initialise un autre dossier sans s’y déplacer.")
    add_code(doc, "qa test")
    add_p(doc, "Lance tous les contrôles du dossier courant.")
    add_code(doc, 'qa test --cwd "C:\\chemin\\vers\\un-autre-projet"')
    add_p(doc, "Teste un autre dossier sans s’y déplacer.")
    add_code(doc, "qa test --fail-fast")
    add_p(doc, "S’arrête au premier contrôle obligatoire en échec.")

    add_h(doc, "9. Problèmes fréquents", 1)
    add_p(doc, "« qa n’est pas reconnu »")
    add_bullet(doc, "Refais npm run build puis npm link dans qa_framework.")
    add_bullet(doc, "Ferme et rouvre PowerShell.")
    add_p(doc, "« qa.config.json introuvable »")
    add_bullet(doc, "Tu n’es pas dans le bon dossier, ou tu n’as pas encore lancé qa init.")
    add_p(doc, "« package.json introuvable »")
    add_bullet(doc, "Lance d’abord npm init -y (ou crée le projet avec Vite), puis qa init.")
    add_p(doc, "« UNABLE_TO_VERIFY_LEAF_SIGNATURE »")
    add_bullet(doc, "Exécute $env:NODE_OPTIONS='--use-system-ca' puis relance.")
    add_p(doc, "« Missing script: lint » (ou typecheck, test:unit…)")
    add_bullet(doc, "Lance qa init, ou désactive le check dans qa.config.json.")
    add_p(doc, "Le rapport ne s’ouvre pas")
    add_bullet(doc, "Ctrl + clic sur le lien file:// affiché, ou ouvre projet\\.qa\\report.html.")
    add_p(doc, "Les tests visuels échouent après un changement d’écran")
    add_bullet(doc, "C’est normal. Si le changement est voulu : npm run test:visual -- --update-snapshots")

    add_h(doc, "10. Fiche mémo (à garder sous la main)", 1)
    add_p(doc, "Première installation de l’outil :")
    add_code(
        doc,
        'cd "...\\qa_framework"\n'
        "$env:NODE_OPTIONS='--use-system-ca'\n"
        "npm install\n"
        "npm run build\n"
        "npm link",
    )
    add_p(doc, "Nouveau projet :")
    add_code(
        doc,
        'cd "...\\mon-projet"\n'
        "npm init -y\n"
        "$env:NODE_OPTIONS='--use-system-ca'\n"
        "qa init\n"
        "qa test",
    )
    add_p(doc, "Projet déjà préparé :")
    add_code(
        doc,
        'cd "...\\mon-projet"\n'
        "$env:NODE_OPTIONS='--use-system-ca'\n"
        "qa test",
    )
    add_note(
        doc,
        "Tu travailles toujours dans le dossier du projet à tester. qa_framework, c’est seulement "
        "l’outil. pilote-web (ou mon-projet), c’est l’application contrôlée. "
        "qa init n’écrase jamais tes fichiers existants.",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
